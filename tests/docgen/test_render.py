import pytest

from legal_assistant.docgen.render import (
    ArticleBlock,
    Attendee,
    MissingContextError,
    TableBlock,
    TextBlock,
    build_context,
    document_text,
    layout_blocks,
    render_document,
)

SHAKHS_WAHED_SCALARS = {
    "company_name": "شركة النور للتوزيع",
    "law_number": "159",
    "law_year": "1981",
    "commercial_registration_no": "303907",
    "commercial_registration_date": "2021/03/14",
    "day_name": "الأحد",
    "day_date": "2026/09/06",
    "names_of_commissioners": "أحمد كامل",
    "owner_name": "مايكل فوزى",
    "company_address": "٤٠ شارع الهرم، الجيزة",
}

ZMM_SCALARS = {
    "company_name": "إنجاز للتجارة",
    "law_number": "159",
    "law_year": "1981",
    "commercial_registration_no": "303907",
    "commercial_registry_office": "استثمار الجيزة",
    "day_name": "الأحد",
    "day_date": "2026/09/06",
    "names_of_commissioners": "أحمد كامل",
    "chairman_name": "مايكل فوزى",
    "chairman_title": "السيد",
    "auditor_name": "الأستاذ/ كمال حسن",
    "secretary_name": "هاني نبيل",
    "vote_counter_1": "رامي عادل",
    "vote_counter_2": "منى سمير",
    "attendance_percentage": "100%",
    "approval_percentage": "100%",
    "company_address": "٤٠ شارع الهرم، الجيزة",
}


def _zmm_articles():
    return [
        ArticleBlock("المادة السادسة", "النص القديم للمادة السادسة.", "النص الجديد.", "المادة (6)"),
        ArticleBlock("المادة السابعة", "النص القديم للمادة السابعة.", "نص جديد آخر.", "المادة (7)"),
    ]


def _articles():
    return [
        ArticleBlock("المادة السادسة", "النص القديم للمادة السادسة.", "النص الجديد."),
        ArticleBlock("المادة السابعة", "النص القديم للمادة السابعة.", "نص جديد آخر."),
    ]


def test_build_context_computes_articles_title_from_the_numbers():
    context = build_context(
        "shakhs_wahed",
        scalars=SHAKHS_WAHED_SCALARS,
        articles=_articles(),
        article_numbers=[6, 7],
        attendees=[],
    )
    assert context["articles_title"] == "المواد السادسة والسابعة"
    assert len(context["articles"]) == 2


def test_render_shakhs_wahed_leaves_no_placeholder_residue():
    context = build_context(
        "shakhs_wahed",
        scalars=SHAKHS_WAHED_SCALARS,
        articles=_articles(),
        article_numbers=[6, 7],
        attendees=[],
    )
    text = document_text(render_document("shakhs_wahed", context))
    assert "{{" not in text and "{%" not in text
    assert "شركة النور للتوزيع" in text
    assert "المواد السادسة والسابعة" in text


def test_render_repeats_the_block_once_per_article():
    context = build_context(
        "shakhs_wahed",
        scalars=SHAKHS_WAHED_SCALARS,
        articles=_articles(),
        article_numbers=[6, 7],
        attendees=[],
    )
    text = document_text(render_document("shakhs_wahed", context))
    assert "النص القديم للمادة السادسة." in text
    assert "النص القديم للمادة السابعة." in text
    assert text.count("قبل التعديل") == 2
    assert text.count("بعد التعديل") == 2


def test_render_a_single_article_uses_the_singular_heading():
    context = build_context(
        "shakhs_wahed",
        scalars=SHAKHS_WAHED_SCALARS,
        articles=_articles()[:1],
        article_numbers=[6],
        attendees=[],
    )
    text = document_text(render_document("shakhs_wahed", context))
    assert "المادة السادسة" in text
    assert "المواد" not in text


def test_render_zmm_lists_each_attendee_with_title_and_role():
    attendees = [
        Attendee("مايكل فوزى", "٩٠", "٩٠", title="السيد", role="مدير الشركة"),
        Attendee("سارة مجدى", "١٠", "١٠", title="السيدة", role="شريك"),
    ]
    context = build_context(
        "zmm",
        scalars=ZMM_SCALARS,
        articles=_zmm_articles()[:1],
        article_numbers=[6],
        attendees=attendees,
    )
    text = document_text(render_document("zmm", context))
    assert "السيد/ مايكل فوزى\n(مدير الشركة)" in text
    assert "السيدة/ سارة مجدى\n(شريك)" in text
    assert "{{" not in text and "{%" not in text


def test_render_zmm_numbers_each_article_as_a_decision_then_the_authorization():
    context = build_context(
        "zmm",
        scalars=ZMM_SCALARS,
        articles=_zmm_articles(),
        article_numbers=[6, 7],
        attendees=[Attendee("مايكل فوزى", title="السيد", role="مدير الشركة")],
    )
    text = document_text(render_document("zmm", context))
    assert "أولاً: الموافقة على تعديل المادة (6) من النظام الأساسي للشركة" in text
    assert "ثانياً: الموافقة على تعديل المادة (7) من النظام الأساسي للشركة" in text
    assert "المادة (7) قبل التعديل:" in text and "المادة (7) بعد التعديل:" in text
    assert "ثالثاً: تفويض كلٍّ من السادة/ أحمد كامل" in text
    assert "سجل تجاري رقم 303907 استثمار الجيزة" in text
    assert "حيث بلغت نسبة حضور الشركاء 100%." in text
    assert "الأستاذ/ كمال حسن\n(مراقب الحسابات)" in text
    assert "بصفته مدير شركة/ إنجاز للتجارة شركة ذات مسئولية محدودة" in text


def test_render_zmm_writes_the_feminine_capacity_for_a_woman_chairing():
    scalars = {**ZMM_SCALARS, "chairman_name": "سارة مجدى", "chairman_title": "السيدة"}
    context = build_context(
        "zmm",
        scalars=scalars,
        articles=_zmm_articles()[:1],
        article_numbers=[6],
        attendees=[Attendee("سارة مجدى", title="السيدة", role="مدير الشركة")],
    )
    text = document_text(render_document("zmm", context))
    assert "دعوة السيدة/ سارة مجدى بصفتها مدير شركة/" in text
    assert "بصفته " not in text


def test_render_raises_when_a_declared_scalar_is_missing():
    scalars = dict(SHAKHS_WAHED_SCALARS)
    del scalars["owner_name"]
    context = build_context(
        "shakhs_wahed",
        scalars=scalars,
        articles=_articles()[:1],
        article_numbers=[6],
        attendees=[],
    )
    with pytest.raises(MissingContextError) as excinfo:
        render_document("shakhs_wahed", context)
    assert "owner_name" in str(excinfo.value)


def test_render_raises_for_an_unknown_company_type():
    with pytest.raises(KeyError):
        render_document("tawsiya_bil_ashum", {})


def test_build_context_requires_at_least_one_article():
    with pytest.raises(ValueError):
        build_context(
            "shakhs_wahed",
            scalars=SHAKHS_WAHED_SCALARS,
            articles=[],
            article_numbers=[],
            attendees=[],
        )


def test_render_document_fails_closed_on_missing_articles_even_bypassing_build_context():
    # render_document is a documented public entry point in its own right;
    # a caller must not be able to skip build_context's article check by
    # constructing the context directly with an empty/missing articles list.
    # articles_title is filled in deliberately so the only gap under test
    # is the empty articles list itself, not an unrelated scalar.
    context = {
        **SHAKHS_WAHED_SCALARS,
        "articles_title": "المادة السادسة",
        "articles": [],
        "attendees": [],
    }
    with pytest.raises(MissingContextError) as excinfo:
        render_document("shakhs_wahed", context)
    assert excinfo.value.missing == {"articles"}


def test_render_document_fails_closed_when_articles_key_is_absent():
    context = {
        **SHAKHS_WAHED_SCALARS,
        "articles_title": "المادة السادسة",
        "attendees": [],
    }
    with pytest.raises(MissingContextError) as excinfo:
        render_document("shakhs_wahed", context)
    assert excinfo.value.missing == {"articles"}


def test_render_document_fails_closed_on_missing_attendees_for_zmm():
    context = build_context(
        "zmm",
        scalars=ZMM_SCALARS,
        articles=_zmm_articles()[:1],
        article_numbers=[6],
        attendees=[],
    )
    with pytest.raises(MissingContextError) as excinfo:
        render_document("zmm", context)
    assert excinfo.value.missing == {"attendees"}


def test_render_document_does_not_require_attendees_for_shakhs_wahed():
    # شخص واحد declares no attendee_placeholders at all, so an empty
    # attendees list must not be treated as a missing value for it.
    context = build_context(
        "shakhs_wahed",
        scalars=SHAKHS_WAHED_SCALARS,
        articles=_articles()[:1],
        article_numbers=[6],
        attendees=[],
    )
    text = document_text(render_document("shakhs_wahed", context))
    assert "{{" not in text and "{%" not in text


def test_render_zmm_puts_company_address_in_the_footer_with_no_stray_braces():
    attendees = [Attendee("مايكل فوزى", "١٠٠", "١٠٠", title="السيد", role="مدير الشركة")]
    context = build_context(
        "zmm",
        scalars=ZMM_SCALARS,
        articles=_zmm_articles()[:1],
        article_numbers=[6],
        attendees=attendees,
    )
    rendered = render_document("zmm", context)
    text = document_text(rendered)
    assert "٤٠ شارع الهرم، الجيزة" in text
    assert "{company_address}" not in text
    assert "{{" not in text and "{%" not in text


# --- layout of article text -------------------------------------------------


_OCR_ARTICLE = (
    "حدد رأس مال الشركة بمبلغ ١٠٠٠٠٠ جنيه مصري، وجميعها\n"
    "حصص نقدية، وقد تم توزيع هذه الحصص بين الشركاء على الوجه الآتي :\n"
    "| م | الاسم | عدد الحصص |\n"
    "|---|---|---|\n"
    "| ١ | أحمد محمود سالم | ٩٠ |\n"
    "| ٢ | سارة علي حسن | ١٠ |\n"
    "وتبلغ نسبة المشاركة المصرية ١٠٠%"
)


def test_layout_joins_scan_line_wraps_and_builds_the_table():
    assert layout_blocks(_OCR_ARTICLE) == [
        TextBlock(
            "حدد رأس مال الشركة بمبلغ ١٠٠٠٠٠ جنيه مصري، وجميعها "
            "حصص نقدية، وقد تم توزيع هذه الحصص بين الشركاء على الوجه الآتي :"
        ),
        TableBlock([
            ["م", "الاسم", "عدد الحصص"],
            ["١", "أحمد محمود سالم", "٩٠"],
            ["٢", "سارة علي حسن", "١٠"],
        ]),
        TextBlock("وتبلغ نسبة المشاركة المصرية ١٠٠%"),
    ]


def test_layout_keeps_every_word_and_only_changes_whitespace():
    blocks = layout_blocks(_OCR_ARTICLE)
    words = []
    for block in blocks:
        if isinstance(block, TextBlock):
            words += block.text.split()
        else:
            words += [w for row in block.rows for cell in row for w in cell.split()]
    source = [w for w in _OCR_ARTICLE.split() if set(w) - set("|-")]
    assert words == source


def test_layout_breaks_paragraphs_at_blank_lines_and_list_items():
    text = "الغرض هو:\n• التجارة العامة\n• التصدير\n\nوذلك دون الإخلال\nبأحكام القانون."
    assert layout_blocks(text) == [
        TextBlock("الغرض هو:"),
        TextBlock("• التجارة العامة"),
        TextBlock("• التصدير"),
        TextBlock("وذلك دون الإخلال بأحكام القانون."),
    ]


def test_layout_pads_ragged_table_rows():
    blocks = layout_blocks("| أ | ب | ج |\n| ١ | ٢ |")
    assert blocks == [TableBlock([["أ", "ب", "ج"], ["١", "٢", ""]])]


def _rendered_shakhs_wahed(original: str) -> bytes:
    context = build_context(
        "shakhs_wahed",
        scalars=SHAKHS_WAHED_SCALARS,
        articles=[ArticleBlock("المادة السادسة", original, "النص الجديد.")],
        article_numbers=[6],
        attendees=[],
    )
    return render_document("shakhs_wahed", context)


def test_render_turns_a_markdown_table_into_a_right_to_left_word_table():
    import io

    import docx
    from docx.oxml.ns import qn

    document = docx.Document(io.BytesIO(_rendered_shakhs_wahed(_OCR_ARTICLE)))
    tables = [t for t in document.tables if t.rows[0].cells[0].text == "م"]
    assert len(tables) == 1
    table = tables[0]
    assert [c.text for c in table.rows[1].cells] == ["١", "أحمد محمود سالم", "٩٠"]
    assert table._tbl.tblPr.find(qn("w:bidiVisual")) is not None
    text = document_text(_rendered_shakhs_wahed(_OCR_ARTICLE))
    assert "|" not in text
    assert "DOCGEN-" not in text


def test_render_leaves_no_paragraph_justified():
    import io
    import zipfile

    data = _rendered_shakhs_wahed(_OCR_ARTICLE)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                assert 'w:jc w:val="both"' not in archive.read(name).decode("utf-8"), name


def test_render_keeps_the_scan_paragraph_as_one_paragraph():
    import io

    import docx

    document = docx.Document(io.BytesIO(_rendered_shakhs_wahed(_OCR_ARTICLE)))
    texts = [p.text for p in document.paragraphs]
    assert (
        "حدد رأس مال الشركة بمبلغ ١٠٠٠٠٠ جنيه مصري، وجميعها "
        "حصص نقدية، وقد تم توزيع هذه الحصص بين الشركاء على الوجه الآتي :"
    ) in texts


def test_column_widths_follow_content_and_fill_the_width():
    from legal_assistant.docgen.render import column_widths

    rows = [["م", "الاسم والجنسية", "عدد الحصص"], ["١", "أحمد محمود سالم", "٩٠"]]
    widths = column_widths(rows, 9000)
    assert widths[1] > widths[2] > widths[0]
    assert 9000 - len(widths) <= sum(widths) <= 9000


def test_generated_paragraph_properties_keep_schema_order():
    import io

    import docx
    from docx.oxml.ns import qn

    order = ["pStyle", "keepNext", "keepLines", "widowControl", "bidi", "spacing", "ind",
             "contextualSpacing", "jc", "rPr"]
    document = docx.Document(io.BytesIO(_rendered_shakhs_wahed(_OCR_ARTICLE)))
    for ppr in document.element.body.iter(qn("w:pPr")):
        seen = [c.tag.split("}")[1] for c in ppr if c.tag.split("}")[1] in order]
        assert seen == sorted(seen, key=order.index), seen


def test_layout_does_not_split_a_pipe_cell_at_a_double_space():
    text = "| م | الاسم | عدد الحصص |\n|---|---|---|\n| ١ | أحمد  محمود سالم | ٩٠ |"
    blocks = layout_blocks(text)
    assert blocks == [TableBlock([["م", "الاسم", "عدد الحصص"], ["١", "أحمد  محمود سالم", "٩٠"]])]


def test_layout_keeps_a_row_of_dashes_as_content():
    blocks = layout_blocks("| البند | القيمة |\n|:---|---:|\n| حصص عينية | - |\n| - | — |")
    assert blocks == [
        TableBlock([["البند", "القيمة"], ["حصص عينية", "-"], ["-", "—"]])
    ]


def test_layout_keeps_every_line_break_of_typed_text():
    typed = "الشريك الأول ٥٠ حصة\nالشريك الثاني ٥٠ حصة"
    assert layout_blocks(typed, join_wraps=False) == [
        TextBlock("الشريك الأول ٥٠ حصة"),
        TextBlock("الشريك الثاني ٥٠ حصة"),
    ]


def test_bold_header_run_keeps_schema_order_after_rfonts():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from legal_assistant.docgen.render import _new_paragraph

    anchor = OxmlElement("w:p")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag in ("w:rStyle", "w:rFonts", "w:sz"):
        rpr.append(OxmlElement(tag))
    run.append(rpr)
    anchor.append(run)

    paragraph = _new_paragraph(anchor, "الاسم", bold=True)
    tags = [e.tag for e in paragraph.find(qn("w:r")).find(qn("w:rPr"))]
    assert tags == [qn(t) for t in ("w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:sz")]
