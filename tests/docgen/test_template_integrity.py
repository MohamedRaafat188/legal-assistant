"""Guard against a silent hand-edit dropping legal prose from a template.

`registry.verify_all()` only checks that the *placeholders* (`{{ ... }}` /
`{% ... %}`) in each shipped template match the declared contract. It says
nothing about the boilerplate sentences sitting *between* those
placeholders. A future hand-edit to a `.docx` template that deleted a
paragraph of legal text, or reshaped a table, would pass `verify_all()` and
the whole rest of the test suite without a single assertion failing.

This file pins, per template: the exact count of non-empty paragraphs, the
exact number of tables, each table's row/column shape, and a SHA-256 digest
of every non-empty paragraph and table-cell string with the docxtpl tokens
(`{{...}}` and `{%...%}`) stripped out first.

If one of these assertions fails:
  - If you (a human) just edited that template on purpose, recompute the
    counts and digest for that template (see the module-level `_measure`
    helper below -- run it against the changed file) and update the pinned
    constant here to match. That is the expected workflow for an intentional
    template change.
  - If you did NOT touch the template, STOP. Something -- a bad merge, a
    corrupted save, an automated tool -- silently altered a legal document
    template. Find out what changed and why before touching this test.
"""

from __future__ import annotations

import hashlib
import re

import docx

from legal_assistant.docgen.templates.registry import CompanyType, get_template

_TOKEN = re.compile(r"\{\{.*?\}\}|\{%.*?%\}", re.DOTALL)

# company_type -> (non_empty_paragraph_count, [(rows, cols), ...], sha256_hex_digest)
_EXPECTED = {
    "shakhs_wahed": (19, [], "2f9aba843b84a860638506cef443a6f007315023d00b6a77faf47f3d2e893a23"),
    "zmm": (
        17,
        [(4, 4), (1, 2)],
        "ee36e75824fdab1f59a710c4a978bb58af557a7652972bbc4340537ad33d6695",
    ),
    "masahma": (
        23,
        [(4, 4), (1, 3)],
        "e9ed04060e3d31bb504884a33652d753c5997cb5c67a8d5e30045b045e66ee0b",
    ),
}


def _header_footer_parts(document):
    """Every distinct header/footer object on a document's sections.

    Mirrors `render.document_text`'s helper of the same purpose: covers the
    default, first-page, and even-page variants, skipping any variant that
    is merely linked to the previous section (i.e. not actually defined) so
    its inherited-empty text is not double-counted.
    """
    parts = []
    for section in document.sections:
        for attr in (
            "header",
            "footer",
            "first_page_header",
            "first_page_footer",
            "even_page_header",
            "even_page_footer",
        ):
            part = getattr(section, attr)
            if not part.is_linked_to_previous:
                parts.append(part)
    return parts


def _measure(path):
    """(non_empty_paragraph_count, [(rows, cols), ...], sha256_hex_digest) for a .docx.

    Counts and hashes headers/footers along with the body, so a hand-edit
    that silently drops or alters boilerplate text sitting in a footer (as
    once shipped verbatim in the zmm/masahma templates -- see task-10
    review finding 3) is caught by this guard too.
    """
    document = docx.Document(str(path))
    paragraph_texts = [p.text for p in document.paragraphs if p.text.strip()]

    shapes = []
    strings = list(paragraph_texts)
    for table in document.tables:
        shapes.append((len(table.rows), len(table.columns)))
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    strings.append(cell.text)

    for hf in _header_footer_parts(document):
        for p in hf.paragraphs:
            if p.text.strip():
                paragraph_texts.append(p.text)
                strings.append(p.text)
        for table in hf.tables:
            shapes.append((len(table.rows), len(table.columns)))
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        strings.append(cell.text)

    cleaned = [_TOKEN.sub("", s) for s in strings]
    digest = hashlib.sha256("".join(cleaned).encode("utf-8")).hexdigest()
    return len(paragraph_texts), shapes, digest


def test_template_prose_matches_the_pinned_content_digest():
    for company_type in CompanyType:
        spec = get_template(company_type.value)
        expected_paragraphs, expected_shapes, expected_digest = _EXPECTED[company_type.value]

        paragraph_count, table_shapes, digest = _measure(spec.path)

        assert paragraph_count == expected_paragraphs, (
            f"{company_type.value}: non-empty paragraph count changed "
            f"({paragraph_count} != {expected_paragraphs}) -- see module docstring"
        )
        assert table_shapes == expected_shapes, (
            f"{company_type.value}: table count/shape changed "
            f"({table_shapes} != {expected_shapes}) -- see module docstring"
        )
        assert digest == expected_digest, (
            f"{company_type.value}: template prose digest changed -- see module docstring"
        )
