"""Fill a company-type template and return .docx bytes.

Nothing in this module writes legal prose. It moves already-approved strings
into a Word template and refuses to emit a file if any declared placeholder
is unfilled -- a document filed at GAFI containing a literal
"{{ owner_name }}" is worse than a 500.
"""

from __future__ import annotations

import copy
import io
import re
import secrets
from collections.abc import Sequence
from dataclasses import asdict, dataclass

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

from legal_assistant.docgen.numbering import ArticleRef, articles_title, decision_ordinal
from legal_assistant.docgen.parsing.commercial_register import is_separator_row, table_cells
from legal_assistant.docgen.templates.registry import get_template


@dataclass(frozen=True)
class ArticleBlock:
    """One قبل/بعد pair as it appears in the document."""

    article_name: str
    article_original_content: str
    article_new_content: str
    # "المادة (3)": the ذ.م.م محضر names articles by number, not in words.
    article_label: str = ""


@dataclass(frozen=True)
class Attendee:
    name: str
    shares: str = ""
    percentage: str = ""
    title: str = ""  # السيد / السيدة / السادة, chosen by the lawyer
    role: str = ""  # مدير الشركة / شريك


class MissingContextError(RuntimeError):
    """A placeholder the template declares has no value in the context."""

    def __init__(self, company_type: str, missing: set[str]) -> None:
        super().__init__(
            f"cannot render {company_type}: missing context values {sorted(missing)}"
        )
        self.company_type = company_type
        self.missing = missing


def build_context(
    company_type: str,
    scalars: dict[str, str],
    articles: Sequence[ArticleBlock],
    article_numbers: Sequence[ArticleRef | int],
    attendees: Sequence[Attendee],
) -> dict:
    """Assemble the docxtpl context.

    `articles_title` and the decision ordinals are computed here rather than
    passed in, so the headings can never disagree with the articles actually
    in the document: article i is decision i («أولاً», «ثانياً», ...) and the
    authorization is the decision after the last article.
    """
    if not articles:
        raise ValueError("a قرار/محضر التعديل must amend at least one article")
    return {
        **scalars,
        "articles_title": articles_title(article_numbers),
        "authorization_ordinal": decision_ordinal(len(articles) + 1),
        "articles": [
            {**asdict(a), "ordinal": decision_ordinal(i)} for i, a in enumerate(articles, 1)
        ],
        "attendees": [asdict(p) for p in attendees],
    }


# --- article text layout ----------------------------------------------------
#
# Article text is stored as OCR wrote it: one line per SCAN line, tables as
# markdown pipe rows. Pasted into Word as-is, every scan line becomes a forced
# break and a table becomes a row of "|" characters. `layout_blocks` turns it
# into paragraphs and tables -- whitespace only; no word is added, dropped or
# changed.


@dataclass(frozen=True)
class TextBlock:
    text: str


@dataclass(frozen=True)
class TableBlock:
    rows: list[list[str]]


_LIST_ITEM = re.compile(r"^(?:[•●▪◦*\-–]|[0-9٠-٩]{1,2}\s*[-).]|[أ-ي]\s*[-)])\s*")
_PARAGRAPH_END = (".", ":", "؛", "：")


def _is_table_line(line: str) -> bool:
    return line.startswith("|")


def layout_blocks(text: str, join_wraps: bool = True) -> list[TextBlock | TableBlock]:
    """Paragraphs and tables, in order.

    With `join_wraps` (OCR text), a paragraph ends at a blank line, a table,
    a list item, or a line ending in «.», «:» or «؛»; other line breaks are
    the scan's line wrapping and are joined with a space. Without it (text
    the lawyer typed), every line break is the lawyer's own and is kept.
    """
    blocks: list[TextBlock | TableBlock] = []
    paragraph: list[str] = []
    table: list[list[str]] = []

    def flush_paragraph() -> None:
        if paragraph:
            blocks.append(TextBlock(" ".join(paragraph)))
            paragraph.clear()

    def flush_table() -> None:
        if table:
            width = max(len(row) for row in table)
            blocks.append(TableBlock([row + [""] * (width - len(row)) for row in table]))
            table.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if _is_table_line(line):
            flush_paragraph()
            if not is_separator_row(line):
                table.append(table_cells(line))
            continue
        flush_table()
        if not line:
            flush_paragraph()
            continue
        if _LIST_ITEM.match(line):
            flush_paragraph()
        paragraph.append(line)
        if line.endswith(_PARAGRAPH_END) or not join_wraps:
            flush_paragraph()
    flush_table()
    flush_paragraph()
    return blocks


_JUSTIFIED = {"both", "distribute", "lowKashida", "mediumKashida", "highKashida", "thaiDistribute"}


def _align_start(element) -> None:
    """Drop justified alignment under `element`. With no `w:jc` a paragraph
    aligns to its start edge -- the right, for Arabic. (An explicit "right"
    would not do: Word reads left/right as start/end in a bidi paragraph.)"""
    for jc in list(element.iter(qn("w:jc"))):
        if jc.get(qn("w:val")) in _JUSTIFIED:
            jc.getparent().remove(jc)


# pPr children that must come AFTER w:spacing (schema order).
_AFTER_SPACING = (
    "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _set_spacing(ppr, **values: str) -> None:
    """Set w:spacing attributes, creating the element in its schema slot."""
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        successor = next(
            (child for child in ppr if child.tag in {qn(t) for t in _AFTER_SPACING}), None
        )
        if successor is None:
            ppr.append(spacing)
        else:
            successor.addprevious(spacing)
    for name, value in values.items():
        spacing.set(qn(f"w:{name}"), value)


def _new_paragraph(anchor, text: str, bold: bool = False, in_cell: bool = False):
    """A paragraph in the anchor paragraph's paragraph and run formatting.
    In a table cell the anchor's indent and spacing are dropped, so rows stay
    tight."""
    paragraph = OxmlElement("w:p")
    ppr = anchor.find(qn("w:pPr"))
    if ppr is not None:
        ppr = copy.deepcopy(ppr)
        if in_cell:
            for tag in ("w:ind", "w:spacing", "w:keepNext", "w:keepLines"):
                for element in ppr.findall(qn(tag)):
                    ppr.remove(element)
            _set_spacing(ppr, before="0", after="0")
        paragraph.append(ppr)
    first_run = anchor.find(qn("w:r"))
    rpr = first_run.find(qn("w:rPr")) if first_run is not None else None
    rpr = copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")
    if bold and rpr.find(qn("w:b")) is None:
        for old in rpr.findall(qn("w:bCs")):
            rpr.remove(old)
        # Schema order: rStyle, rFonts, then b, bCs.
        leading = [e for e in rpr if e.tag in {qn("w:rStyle"), qn("w:rFonts")}]
        for element in (OxmlElement("w:bCs"), OxmlElement("w:b")):
            if leading:
                leading[-1].addnext(element)
            else:
                rpr.insert(0, element)
    run = OxmlElement("w:r")
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    paragraph.append(run)
    _align_start(paragraph)
    return paragraph


def column_widths(rows: list[list[str]], total: int) -> list[int]:
    """Split `total` (twips) across the columns in proportion to their
    longest cell, with a floor so a «م» column stays readable."""
    weights = [max(4, *(len(row[i]) for row in rows)) for i in range(len(rows[0]))]
    return [total * w // sum(weights) for w in weights]


def _new_table(anchor, rows: list[list[str]], text_width: int):
    """A bordered, right-to-left table spanning the text width. The first row
    is the header: bold, and repeated when the table crosses a page."""
    widths = column_widths(rows, text_width)
    tbl = OxmlElement("w:tbl")
    # tblPr children in schema order: bidiVisual, tblW, tblBorders, tblLayout.
    tblpr = OxmlElement("w:tblPr")
    tblpr.append(OxmlElement("w:bidiVisual"))  # first column on the right
    tblpr.append(_width("w:tblW", sum(widths)))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    tblpr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    tbl.append(tblpr)
    grid = OxmlElement("w:tblGrid")
    for column in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(column))
        grid.append(col)
    tbl.append(grid)
    for index, cells in enumerate(rows):
        tr = OxmlElement("w:tr")
        if index == 0:
            trpr = OxmlElement("w:trPr")
            trpr.append(OxmlElement("w:tblHeader"))
            tr.append(trpr)
        for column, value in zip(widths, cells, strict=True):
            tc = OxmlElement("w:tc")
            tcpr = OxmlElement("w:tcPr")
            tcpr.append(_width("w:tcW", column))
            tc.append(tcpr)
            tc.append(_new_paragraph(anchor, value, bold=index == 0, in_cell=True))
            tr.append(tc)
        tbl.append(tr)
    return tbl


def _width(tag: str, twips: int):
    element = OxmlElement(tag)
    element.set(qn("w:w"), str(twips))
    element.set(qn("w:type"), "dxa")
    return element


def _expand_marker(anchor, text: str, text_width: int, join_wraps: bool = True) -> None:
    """Replace the marker paragraph `anchor` with `text` laid out as
    paragraphs and tables, in the anchor's formatting."""
    after_table = False
    for block in layout_blocks(text, join_wraps) or [TextBlock("")]:
        if isinstance(block, TableBlock):
            anchor.addprevious(_new_table(anchor, block.rows, text_width))
            after_table = True
            continue
        paragraph = _new_paragraph(anchor, block.text)
        if after_table:
            # Breathing room between a table and the text under it (6pt).
            ppr = paragraph.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                paragraph.insert(0, ppr)
            _set_spacing(ppr, before="120")
            after_table = False
        anchor.addprevious(paragraph)
    anchor.getparent().remove(anchor)


_LAID_OUT = ("article_original_content", "article_new_content")


def render_document(company_type: str, context: dict) -> bytes:
    """Render the template for `company_type` into .docx bytes.

    Fails closed not only on a missing scalar but also on a template whose
    article/attendee loop would render as silently empty: a قرار تعديل with
    zero amendments, or a محضر جمعية عامة with zero attendees, is not a
    valid instrument even though docxtpl would happily render it with no
    `{{`/`{%` residue at all. This check is enforced here -- not only in
    `build_context` -- because `render_document` is itself a public entry
    point a caller could reach with a hand-built context.
    """
    spec = get_template(company_type)

    missing = {name for name in spec.scalar_placeholders if not context.get(name)}
    if spec.article_placeholders and not context.get("articles"):
        missing.add("articles")
    if spec.attendee_placeholders and not context.get("attendees"):
        missing.add("attendees")
    if missing:
        raise MissingContextError(company_type, missing)

    # Article text goes in as a unique marker; each marker paragraph is then
    # replaced by real paragraphs and tables (see `layout_blocks`).
    token = secrets.token_hex(8)
    texts: dict[str, str] = {}
    typed: set[str] = set()  # markers holding lawyer-typed text
    articles = []
    for index, article in enumerate(context.get("articles", [])):
        article = dict(article)
        for key in _LAID_OUT:
            marker = f"DOCGEN-{token}-{index}-{key}"
            texts[marker] = article.get(key) or ""
            if key == "article_new_content":
                typed.add(marker)
            article[key] = marker
        articles.append(article)
    template = DocxTemplate(str(spec.path))
    template.render({**context, "articles": articles})
    rendered = io.BytesIO()
    template.save(rendered)

    document = docx.Document(io.BytesIO(rendered.getvalue()))
    body = document.element.body
    section = document.sections[-1]
    # Twips: python-docx lengths are EMU, 635 EMU per twip.
    text_width = (section.page_width - section.left_margin - section.right_margin) // 635
    for paragraph in list(body.iter(qn("w:p"))):
        text = "".join(t.text or "" for t in paragraph.iter(qn("w:t"))).strip()
        if text in texts:
            _expand_marker(paragraph, texts.pop(text), text_width, join_wraps=text not in typed)
    if texts or token in "".join(t.text or "" for t in body.iter(qn("w:t"))):
        # The template puts article text inside other text: fail closed
        # rather than ship a marker. Content-free message.
        raise RuntimeError(f"cannot lay out article text in the {company_type} template")
    _align_start(body)
    for part in _header_footer_parts(document):
        _align_start(part._element)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _header_footer_parts(document):
    """Every distinct header/footer object on a document's sections.

    Covers the default, first-page, and even-page variants; a variant that
    is merely linked to the previous section (i.e. not actually defined)
    is skipped so its inherited-empty text is not double-counted.
    """
    parts = []
    for section in document.sections:
        for attr in (
            "header",
            "footer",
            "first_page_header",
            "first_page_footer",
            "even_page_header",
            "even_page_footer",
        ):
            part = getattr(section, attr)
            if not part.is_linked_to_previous:
                parts.append(part)
    return parts


def document_text(docx_bytes: bytes) -> str:
    """All paragraph and table text of a .docx, newline-joined.

    Includes section headers and footers (and any tables nested in them),
    not just the body -- a stray unrendered token left in a footer must be
    just as visible to `"{{" not in text`-style assertions as one in the
    body.

    Used by tests and by the validation script to assert on a rendered
    document without opening Word.
    """
    document = docx.Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for hf in _header_footer_parts(document):
        parts.extend(p.text for p in hf.paragraphs)
        for table in hf.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
